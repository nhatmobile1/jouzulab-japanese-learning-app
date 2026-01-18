#!/usr/bin/env python3
"""
Genki Excel Parser
Parses Genki textbook word index (.xlsx) into iOS-ready JSON format.

Usage:
    python parse_genki_excel.py --input genki_3rde_word_index.xlsx --output data/genki_deck.json

Requirements:
    pip install openpyxl
"""

import json
import argparse
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    exit(1)


def parse_genki_excel(input_path: str, output_path: str, deck_id: str = "genki_3rd"):
    """Parse Genki Excel file and output JSON."""

    print(f"Loading Excel file: {input_path}")
    wb = openpyxl.load_workbook(input_path, data_only=True)

    entries = []
    entry_count = 0

    # Process each sheet
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\nProcessing sheet: {sheet_name}")

        # Get headers from first row
        headers = []
        for cell in sheet[1]:
            headers.append(str(cell.value).lower().strip() if cell.value else "")

        print(f"  Headers found: {headers}")

        # Try to identify columns
        japanese_col = None
        reading_col = None
        english_col = None
        lesson_col = None

        for i, header in enumerate(headers):
            header_lower = header.lower()
            if any(x in header_lower for x in ['japanese', 'kanji', '日本語', '漢字', 'word']):
                japanese_col = i
            elif any(x in header_lower for x in ['reading', 'hiragana', 'kana', 'ひらがな', '読み']):
                reading_col = i
            elif any(x in header_lower for x in ['english', 'meaning', 'definition', '英語', '意味']):
                english_col = i
            elif any(x in header_lower for x in ['lesson', 'chapter', 'unit', '課']):
                lesson_col = i

        # If columns not identified, try common patterns
        if japanese_col is None and len(headers) >= 1:
            # First column might be Japanese
            japanese_col = 0
        if reading_col is None and len(headers) >= 2:
            reading_col = 1
        if english_col is None and len(headers) >= 3:
            english_col = 2
        if lesson_col is None and len(headers) >= 4:
            lesson_col = 3

        print(f"  Column mapping: japanese={japanese_col}, reading={reading_col}, english={english_col}, lesson={lesson_col}")

        # Process rows (skip header)
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # Skip empty rows
            if not row or all(cell is None or str(cell).strip() == '' for cell in row):
                continue

            # Extract values
            japanese = str(row[japanese_col]).strip() if japanese_col is not None and row[japanese_col] else None
            reading = str(row[reading_col]).strip() if reading_col is not None and len(row) > reading_col and row[reading_col] else None
            english = str(row[english_col]).strip() if english_col is not None and len(row) > english_col and row[english_col] else None
            lesson = str(row[lesson_col]).strip() if lesson_col is not None and len(row) > lesson_col and row[lesson_col] else None

            # Skip if no Japanese text
            if not japanese or japanese == 'None':
                continue

            # Clean up reading (remove 'None' strings)
            if reading == 'None':
                reading = None
            if english == 'None':
                english = None
            if lesson == 'None':
                lesson = None

            entry_count += 1
            entry_id = f"{deck_id}_{entry_count:05d}"

            # Detect entry type
            entry_type = detect_entry_type(japanese)

            # Detect JLPT level based on Genki lessons
            jlpt_level = detect_jlpt_from_lesson(lesson)

            entry = {
                "id": entry_id,
                "japanese": japanese,
                "reading": reading,
                "english": english,
                "entry_type": entry_type,
                "lesson_date": None,
                "tags": ["genki", "textbook"],
                "grammar_patterns": [],
                "jlpt_level": jlpt_level,
                "context_note": None,
                "is_sub_entry": False,
                "source_line": row_num,
                "lesson_frequency": 1,
                "lesson": lesson
            }

            entries.append(entry)

    wb.close()

    # Create output structure
    output = {
        "metadata": {
            "version": "1.0",
            "created_date": datetime.now().strftime("%Y-%m-%d"),
            "total_entries": len(entries),
            "entries_with_reading": sum(1 for e in entries if e["reading"]),
            "entries_with_english": sum(1 for e in entries if e["english"]),
            "entries_missing_reading": sum(1 for e in entries if not e["reading"]),
            "entries_missing_english": sum(1 for e in entries if not e["english"]),
            "lesson_count": len(set(e["lesson"] for e in entries if e["lesson"])),
            "source": "Genki 3rd Edition Word Index"
        },
        "entries": entries
    }

    # Write output
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*50}")
    print(f"Parsing complete!")
    print(f"  Total entries: {len(entries)}")
    print(f"  With reading: {output['metadata']['entries_with_reading']}")
    print(f"  With English: {output['metadata']['entries_with_english']}")
    print(f"  Lessons found: {output['metadata']['lesson_count']}")
    print(f"  Output: {output_path}")

    return entries


def detect_entry_type(text: str) -> str:
    """Detect if text is vocab, phrase, or sentence."""
    if '。' in text or '？' in text or '?' in text:
        return "sentence"
    if ' ' in text or '　' in text or len(text) > 10:
        return "phrase"
    return "vocab"


def detect_jlpt_from_lesson(lesson: str) -> str:
    """Estimate JLPT level based on Genki lesson number."""
    if not lesson:
        return None

    # Extract lesson number
    import re
    match = re.search(r'(\d+)', str(lesson))
    if not match:
        return None

    lesson_num = int(match.group(1))

    # Genki I (Lessons 1-12) covers mostly N5
    # Genki II (Lessons 13-23) covers N5-N4
    if lesson_num <= 6:
        return "N5"
    elif lesson_num <= 12:
        return "N5"
    elif lesson_num <= 17:
        return "N4"
    elif lesson_num <= 23:
        return "N4"
    else:
        return "N4"


def main():
    parser = argparse.ArgumentParser(description="Parse Genki Excel word index to JSON")
    parser.add_argument("--input", "-i", required=True, help="Input Excel file (.xlsx)")
    parser.add_argument("--output", "-o", required=True, help="Output JSON file")
    parser.add_argument("--deck-id", "-d", default="genki_3rd", help="Deck ID prefix for entries")

    args = parser.parse_args()

    parse_genki_excel(args.input, args.output, args.deck_id)


if __name__ == "__main__":
    main()
